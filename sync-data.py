#!/usr/bin/env python3
"""Run after editing beers.json (or any app file) and before committing.

- Regenerates beers-data.js, the offline fallback bundle.
- Regenerates the full beer list downloads in downloads/ (.xlsx needs: pip install openpyxl).
- Stamps sw.js with a hash of the app files so phones drop their old cache.
"""
import csv
import io
import hashlib
import json
import re
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit('openpyxl is needed for the Excel download: pip install openpyxl')

root = Path(__file__).parent

beers = (root / 'beers.json').read_text(encoding='utf-8')
(root / 'beers-data.js').write_text('window.BEERS_DATA = ' + beers + ';\n', encoding='utf-8')
print('Updated beers-data.js')


def style_category(style):
    """Mirrors getStyleCategory() in app.js."""
    s = (style or '').lower()
    if not s: return 'Other'
    if 'non-alcoholic' in s: return 'Non-Alcoholic'
    if s.startswith('ipa') or ' ipa' in s: return 'IPA'
    if 'stout' in s or 'porter' in s: return 'Stout & Porter'
    if any(k in s for k in ('sour', 'wild', 'lambic', 'gose', 'berliner')): return 'Sour & Wild'
    if any(k in s for k in ('pilsner', 'lager', 'keller', 'helles', 'dunkel', 'schwarzbier', 'rauchbier')): return 'Lager & Pilsner'
    if s.startswith('pale ale'): return 'Pale Ale'
    if any(k in s for k in ('barley wine', 'barleywine', 'rye wine', 'wheat wine')): return 'Barley Wine'
    if any(k in s for k in ('belgian', 'saison', 'garde', 'tripel', 'dubbel', 'quadrupel')): return 'Belgian & Saison'
    if any(k in s for k in ('wine', 'cider', 'mead', 'seltzer')): return 'Wine & Cider'
    return 'Other Ale'


SESSION_ORDER = {'Fredag': 1, 'Lørdag tidlig': 2, 'Lørdag sen': 3}
HEADERS = ['ID', 'Beer', 'Brewery', 'Country', 'Style', 'Category', 'Session', 'Hall', 'Stand', 'ABV %',
           'Untappd Rating', 'Untappd URL', 'Served By', 'Note']
WIDTHS = [11, 38, 30, 15, 32, 17, 14, 18, 7, 7, 9, 45, 14, 50]

rows = []
for b in json.loads(beers):
    loc = b.get('location') or ''
    stand = re.search(r'#(\d+)', loc)
    rows.append([b['id'], b['name'], b['brewery'], b['country'], b['style'], style_category(b['style']),
                 b['session'], loc.split('#')[0].strip(), int(stand.group(1)) if stand else None, b.get('alc'),
                 b.get('rating') or None, b.get('untappd') or None, b.get('served_by'), b.get('note')])
rows.sort(key=lambda r: (SESSION_ORDER.get(r[6], 9), r[7], r[8] or 999, r[2], r[1]))

downloads = root / 'downloads'
downloads.mkdir(exist_ok=True)
download_name = 'aarhus-together-2026-beers'

csv_buffer = io.StringIO()
writer = csv.writer(csv_buffer)
writer.writerow(HEADERS)
writer.writerows([['' if v is None else v for v in r] for r in rows])
# UTF-8 with BOM so Excel shows æ/ø/å correctly
csv_bytes = csv_buffer.getvalue().encode('utf-8-sig')


def write_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Beers'
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name='Arial', size=10)
    for cell in ws[1]:
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='B45309')
        cell.alignment = Alignment(vertical='center')
    for i in range(2, ws.max_row + 1):
        ws.cell(i, HEADERS.index('ABV %') + 1).number_format = '0.0'
        ws.cell(i, HEADERS.index('Untappd Rating') + 1).number_format = '0.00'
        link = ws.cell(i, HEADERS.index('Untappd URL') + 1)
        if link.value:
            link.hyperlink = link.value
            link.font = Font(name='Arial', size=10, color='0563C1', underline='single')
    for i, width in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


csv_path = downloads / (download_name + '.csv')
xlsx_path = downloads / (download_name + '.xlsx')
# .xlsx files embed their save time, so only rebuild when the data changed to keep git diffs clean
if csv_path.exists() and xlsx_path.exists() and csv_path.read_bytes() == csv_bytes:
    print('downloads/ already up to date')
else:
    csv_path.write_bytes(csv_bytes)
    write_xlsx(xlsx_path)
    print(f'Updated downloads/ ({len(rows)} beers)')

digest = hashlib.sha256()
for name in ('beers.json', 'app.js', 'style.css', 'index.html', 'vendor/fuse.basic.min.js'):
    digest.update((root / name).read_bytes())
version = digest.hexdigest()[:8]

sw = root / 'sw.js'
text = sw.read_text(encoding='utf-8')
new_text, count = re.subn(r"const CACHE_VERSION = '[^']*';", f"const CACHE_VERSION = '{version}';", text)
if count != 1:
    raise SystemExit('Could not find CACHE_VERSION in sw.js')
sw.write_text(new_text, encoding='utf-8')
print('Set sw.js CACHE_VERSION to ' + version)
